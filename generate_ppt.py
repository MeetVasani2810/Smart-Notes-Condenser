"""
generate_ppt.py
---------------
Generates an 8-slide academic presentation (.pptx) for Smart Notes Condenser.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="Smart_Notes_Condenser_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette (Dark Theme / Professional Navy)
    BG_COLOR = RGBColor(15, 17, 30)       # #0f111e
    CARD_BG = RGBColor(28, 33, 58)        # #1c213a
    ACCENT_PURPLE = RGBColor(99, 102, 241)# #6366f1
    ACCENT_CYAN = RGBColor(56, 189, 248)  # #38bdf8
    ACCENT_GREEN = RGBColor(52, 211, 153) # #34d399
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_LIGHT = RGBColor(203, 213, 225)  # #cbd5e1
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_COLOR
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, title_text, category="SMART NOTES CONDENSER · DL/NLP MINI-PROJECT"):
        # Category label
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, title, items, badge_color=ACCENT_PURPLE):
        # Card container
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(60, 68, 110)
        card.line.width = Pt(1)

        # Content
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        if title:
            p_head = tf.paragraphs[0]
            p_head.text = title
            p_head.font.size = Pt(17)
            p_head.font.bold = True
            p_head.font.color.rgb = badge_color
            p_head.space_after = Pt(12)

        for i, item in enumerate(items):
            p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
            p.text = f"•  {item}"
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(8)

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ──────────────────────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Badge
    b1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.66), Inches(1.2), Inches(4.0), Inches(0.45))
    b1.fill.solid()
    b1.fill.fore_color.rgb = CARD_BG
    b1.line.color.rgb = ACCENT_PURPLE
    tf_b1 = b1.text_frame
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "DEEP LEARNING + NLP MINI PROJECT"
    p_b1.alignment = PP_ALIGN.CENTER
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = ACCENT_CYAN

    # Main Title Box
    t1_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.33), Inches(2.2))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "🧠 Smart Notes Condenser"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Automated Study Material Synthesis via Transformer Deep Learning & Hybrid LLM Inference"
    p1_sub.alignment = PP_ALIGN.CENTER
    p1_sub.font.size = Pt(17)
    p1_sub.font.color.rgb = ACCENT_PURPLE
    p1_sub.space_before = Pt(12)

    # Info card at bottom
    info_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(4.4), Inches(9.33), Inches(2.3))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = CARD_BG
    info_card.line.color.rgb = RGBColor(60, 68, 110)

    tb_info = s1.shapes.add_textbox(Inches(2.3), Inches(4.6), Inches(8.73), Inches(1.9))
    tf_info = tb_info.text_frame
    tf_info.word_wrap = True

    p_info1 = tf_info.paragraphs[0]
    p_info1.text = "Core Tech Stack: Google FLAN-T5 (PyTorch CPU) + Groq Cloud LPU + FastAPI + React 19"
    p_info1.font.size = Pt(14)
    p_info1.font.bold = True
    p_info1.font.color.rgb = TEXT_WHITE
    p_info1.space_after = Pt(10)

    p_info2 = tf_info.add_paragraph()
    p_info2.text = "Key Features: Abstractive Summaries · 5 Key Points · Self-Test Quizzes · Active-Recall Flashcards"
    p_info2.font.size = Pt(13)
    p_info2.font.color.rgb = ACCENT_GREEN
    p_info2.space_after = Pt(10)

    p_info3 = tf_info.add_paragraph()
    p_info3.text = "Presented by: Student / Project Team | Faculty Guide: Department of Computer Science & AI"
    p_info3.font.size = Pt(12)
    p_info3.font.color.rgb = TEXT_MUTED

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 2: Problem Statement & Objectives
    # ──────────────────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement & Project Objectives")

    add_card(
        s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "⚠️ The Problem",
        [
            "Information Overload: Students spend excessive hours manually summarizing lengthy lecture transcripts and textbooks.",
            "Revision Fatigue: Converting raw, unstructured notes into active recall flashcards and quiz questions is labor-intensive.",
            "Small Local Model Limitations: Traditional small local models (<100M params) often truncate or hallucinate on multi-part tasks.",
            "Cloud Fragility: Pure cloud API apps fail during network dropouts or quota exhaustion."
        ],
        ACCENT_CYAN
    )

    add_card(
        s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "🎯 Project Objectives",
        [
            "Automate 4 Core Study Assets: Abstractive summaries, key takeaways, self-test quizzes, and interactive flashcards.",
            "Hybrid Multi-Engine Architecture: Integrate ultra-fast Groq LPU with a 100% offline local PyTorch FLAN-T5 fallback.",
            "Zero GPU Requirement for Local Mode: Enable CPU inference with optimized memory usage on standard laptops.",
            "Interactive UI: Build a modern self-test interface with 'Check Answer' reveals and flip flashcards."
        ],
        ACCENT_GREEN
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 3: Deep Learning & NLP Fundamentals
    # ──────────────────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Deep Learning & NLP Technical Foundation")

    add_card(
        s3, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "1. Transformer Attention",
        [
            "Multi-Head Self-Attention: softmax((Q·Kᵀ)/√d)·V captures global token dependencies simultaneously.",
            "Replaced legacy sequential RNNs & LSTMs with parallelizable matrix operations.",
            "Positional encodings preserve order and syntactic hierarchy."
        ],
        ACCENT_PURPLE
    )

    add_card(
        s3, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "2. Google FLAN-T5-Small",
        [
            "Encoder-Decoder Architecture: Bidirectional encoder + Autoregressive decoder.",
            "Instruction-Tuned on 1800+ NLP tasks treating every problem as text-to-text.",
            "Lightweight footprint (~80M parameters, ~300MB), operating inside ~1.2 GB RAM."
        ],
        ACCENT_CYAN
    )

    add_card(
        s3, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "3. NLP Task Spectrum",
        [
            "Abstractive Summarization: Generates novel paraphrased sentences.",
            "Key Point Extraction: Semantic segmentation into distinct bullet facts.",
            "Question Generation (QAG): Synthesizes self-testing Q&A pairs.",
            "Knowledge Extraction: Dual-sided active-recall memory cards."
        ],
        ACCENT_GREEN
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 4: System Architecture & Hybrid Engine
    # ──────────────────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "System Architecture & Hybrid Dual-Engine Pipeline")

    add_card(
        s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "⚡ Primary Engine: Groq LPU Cloud",
        [
            "Hardware: Language Processing Unit (LPU) designed for deterministic ultra-high throughput.",
            "Model: openai/gpt-oss-120b & llama-3.3-70b-versatile.",
            "Sub-Second Latency: Typical inference generation in 0.7 – 1.0s.",
            "Dynamic Diversity: Temperature tuning (T=0.85) and Focus Angle nonces for fresh quiz generation."
        ],
        ACCENT_PURPLE
    )

    add_card(
        s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "💻 Fallback Engine: Local PyTorch FLAN-T5",
        [
            "100% Offline Capability: Operates seamlessly without internet or external API dependencies.",
            "PyTorch Inference: Executed with torch.no_grad() and model.eval() on CPU.",
            "Beam Search Decoding: num_beams=4, length_penalty=1.2, no_repeat_ngram_size=3.",
            "Automatic Failover: Seamlessly activates if API limits or network outages occur."
        ],
        ACCENT_CYAN
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 5: Core Features & Study Formats
    # ──────────────────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Core Features & Interactive Student User Experience")

    add_card(
        s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.4),
        "📄 Abstractive Summary",
        [
            "Academic left-accent card presentation.",
            "Cohesive 1-2 paragraph overview synthesizing core concepts, definitions, and equations."
        ],
        ACCENT_PURPLE
    )

    add_card(
        s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.4),
        "🔑 5 Key Points",
        [
            "Structured numbered cards (1 through 5).",
            "Eliminates run-on text; ensures each bullet represents a distinct takeaway."
        ],
        ACCENT_CYAN
    )

    add_card(
        s5, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.4),
        "❓ Interactive Self-Test Quiz",
        [
            "Questions render first; answers stay hidden.",
            "Interactive '🔍 Check Answer' / '🙈 Hide Answer' buttons with 'Reveal All' toggle."
        ],
        ACCENT_GREEN
    )

    add_card(
        s5, Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.4),
        "🃏 3D Study Flashcards",
        [
            "Interactive flip cards testing terms and definitions.",
            "Click-to-flip Front/Back cards for active-recall memory retention."
        ],
        RGBColor(244, 114, 182)
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 6: Technical Implementation & Innovations
    # ──────────────────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Technical Implementation & Engineering Innovations")

    add_card(
        s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "🧠 AI & Prompt Engineering",
        [
            "Focus-Angle Rotation: Randomizes testing perspectives (mechanisms, equations, terminology, cause-effect) on every click.",
            "Dual-Call Summary Synthesis: Local model merges overview definitions with specific factor/stage calls.",
            "Section Chunking for Local Bullets: Notes are partitioned into 2-sentence micro-sections to prevent repetitive bullets.",
            "Pydantic Validation: Strict schema sanitization up to 20,000 characters."
        ],
        ACCENT_PURPLE
    )

    add_card(
        s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "⚛️ Full-Stack Web Architecture",
        [
            "FastAPI Lifespan Caching: Preloads model weights into RAM on server boot, eliminating per-request loading delay.",
            "React 19 + Vite 8: Modern single-page UI built with sub-second HMR and 1.3s production build.",
            "Live Dual Counters: Live tracking of word count and character count (≤ 20,000 chars).",
            "1-Click Clipboard Copy: Instant copying with dynamic '✓ Copied!' toast feedback."
        ],
        ACCENT_GREEN
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 7: Performance Evaluation & Benchmarks
    # ──────────────────────────────────────────────────────────────────────────
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Performance Evaluation & Experimental Results")

    add_card(
        s7, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0),
        "📊 Comparative Evaluation Metrics",
        [
            "Inference Speed: Groq Cloud LPU generates full study sets in 0.7 – 1.0s (40x faster than CPU inference).",
            "Local Model Efficiency: FLAN-T5-Small operates stably inside ~1.2 GB RAM with 0 GPU requirement.",
            "Quiz Diversity Test: Generating on the same text 5 times produced 5 completely distinct sets of exam questions.",
            "Multi-Domain Robustness: Validated across Biology (Photosynthesis), Computer Science (Binary Search), and Physics.",
            "Zero Hallucination Guard: Strict system prompts restrict content derivation 100% to the input notes."
        ],
        ACCENT_CYAN
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SLIDE 8: Conclusion & Future Scope
    # ──────────────────────────────────────────────────────────────────────────
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Conclusion & Future Roadmap")

    add_card(
        s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "🏁 Project Conclusions",
        [
            "Successfully engineered a complete, production-ready educational AI web application.",
            "Demonstrated practical applications of Encoder-Decoder Transformer models and modern LLM pipelines.",
            "Balanced high speed (Groq LPU) with academic robustness (offline PyTorch CPU fallback).",
            "Delivered an engaging self-test user experience designed specifically for students."
        ],
        ACCENT_GREEN
    )

    add_card(
        s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "🚀 Future Enhancements",
        [
            "Multi-Format Document Upload: Direct ingestion of PDF, PPTX, and DOCX files.",
            "OCR for Handwritten Notes: Optical character recognition for digitized student notebooks.",
            "Multilingual NLP: Summarization and translation in regional languages.",
            "Audio Podcast Generation: Text-to-speech synthesis for auditory revision."
        ],
        ACCENT_PURPLE
    )

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Smart_Notes_Condenser_Presentation.pptx"
    create_presentation(out)
